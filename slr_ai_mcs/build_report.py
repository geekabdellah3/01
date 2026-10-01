# -*- coding: utf-8 -*-
import os
from collections import Counter
from data import F, ART, BLOCKS, NOT_EXTRACTABLE, SECONDARY, EXCLUDED
from build_stats import (REL, COND, ALLREL, evidence, rel_row, fmt, arts, block_stats)
import narrative as N

OUT = os.path.dirname(os.path.abspath(__file__))
ev = evidence()
bs = block_stats()


def esc(s):
    return str(s).replace("|", "/").replace("\n", " ")


def table(header, rows):
    out = ["| " + " | ".join(header) + " |", "|" + "|".join(["---"] * len(header)) + "|"]
    for r in rows:
        out.append("| " + " | ".join(esc(c) for c in r) + " |")
    return "\n".join(out)


def num(v):
    return {0.0: "0", 0.5: "0.5", 1.0: "1"}[v]


def art_label(a):
    return f"{a} ({ART[a]['short']})"

# ---------------------------------------------------------------- corpus triage rows
DESIGN = {"Q": "Quantitative", "M": "Mixed", "L": "Qualitative", "S": "Secondary / review / conceptual", "X": "Excluded (vendor discourse)"}
tri_rows = []
for a, m in ART.items():
    if a in NOT_EXTRACTABLE:
        use = "No extractable findings"
    elif m["d"] == "S":
        use = "Review-level corroboration only"
    elif m["d"] == "X":
        use = "Excluded from counts"
    else:
        n = sum(1 for x in F if x["art"] == a)
        use = f"{n} finding(s) coded"
    tri_rows.append([a, m["short"], m["y"], DESIGN[m["d"]], m["method"], m["sample"], m["ctx"], use, m["note"]])

cnt = Counter(m["d"] for m in ART.values())
n_quant = cnt["Q"] + cnt["M"]
n_qual_ok = sum(1 for a, m in ART.items() if m["d"] == "L" and a not in NOT_EXTRACTABLE)

# ---------------------------------------------------------------- findings table
find_hdr = ["ID", "Article", "Empirical finding", "Independent variable / antecedent", "Mechanism", "Mediator",
            "Moderator / condition", "Outcome", "Direction", "Significant?", "Method", "Sample", "Context"]
find_rows = []
for x in F:
    m = ART[x["art"]]
    find_rows.append([x["id"], x["art"], x["finding"], x["iv"], x["mech"], x["med"], x["mod"], x["out"], x["dir"], x["sig"],
                      m["method"], m["sample"], m["ctx"]])

code_hdr = ["ID", "Article", "Empirical finding (short)", "Sensing", "Sensemaking", "Business Partner", "Translation", "Planning",
            "Cybernetic", "Administrative", "Cultural", "Rewards", "Conditions", "Outcomes"]
code_rows = [[x["id"], x["art"], x["short"]] + [num(v) for v in x["codes"]] for x in F]

# ---------------------------------------------------------------- relationship rows
rel_hdr = ["Relationship", "Supporting empirical studies", "Number of studies", "Quantitative evidence", "Qualitative evidence",
           "Null / contradictory evidence", "Reading"]

def rel_table(keys, labels):
    rows = []
    for k in keys:
        r = rel_row(k, ev)
        sup_qd = arts(r["supQ"] + r["supD"]); sup_l = arts(r["supL"])
        ind = arts(r["indQ"] + r["indL"]); nul = arts(r["nul"]); unv = arts(r["unv"])
        sup_all = sorted(set(sup_qd) | set(sup_l), key=lambda s: int(s[1:]))
        support = ", ".join(sup_all) if sup_all else "none"
        if ind:
            support += f" | indirect: {', '.join(ind)}"
        if unv:
            support += f" | abstract-only (unverified): {', '.join(unv)}"
        n = f"{len(sup_all)} direct ({len(sup_qd)} quant, {len(sup_l)} qual)"
        if ind:
            n += f" + {len(ind)} indirect"
        if nul:
            n += f"; {len(nul)} null/contradictory"
        quant = fmt(r["supQ"] + r["supD"])
        if r["indQ"]:
            quant += " [indirect: " + fmt(r["indQ"]) + "]"
        if unv:
            quant += " [abstract-only: " + fmt(r["unv"]) + "]" if quant != "-" else "[abstract-only: " + fmt(r["unv"]) + "]"
        qual = fmt(r["supL"])
        if r["indL"]:
            qual += (" [indirect: " + fmt(r["indL"]) + "]") if qual != "-" else "[indirect: " + fmt(r["indL"]) + "]"
        nullc = fmt(r["nul"])
        rows.append([labels[k], support, n, quant, qual, nullc, N.REL_READING.get(k, "")])
    return rows

req_rows = rel_table([k for k in REL if k.startswith("R")], REL)
em_rows = rel_table([k for k in REL if k.startswith("E")], REL)
cond_rows = rel_table(list(COND), COND)

# ---------------------------------------------------------------- block stats table
bst_rows = []
for b in BLOCKS:
    s = bs[b]
    bst_rows.append([b, s["f1"], s["fh"], f"{len(s['aq'])}: " + ", ".join(s["aq"]) if s["aq"] else "0",
                     f"{len(s['al'])}: " + ", ".join(s["al"]) if s["al"] else "0"])
bst_hdr = ["Block", "Findings coded 1", "Findings coded 0.5", "Quantitative articles with a finding coded 1", "Qualitative articles with a finding coded 1"]

# ---------------------------------------------------------------- secondary table
SEC = [
 ("A30", "Review (Guenther-based): routine tasks shift to BP / data-scientist / governance profiles; instruments adapt; decision-making decentralises; MC integrates across functions; information asymmetry falls; anxiety/aversion possible; analytic and soft skills needed."),
 ("A38", "Review 2007-2017: power shifts horizontally to those who master digital technology; accountants risk losing legitimacy if they stay out of digital initiatives; names two gaps - the effect of digitalisation on financial performance and the role of contextual variables."),
 ("A20", "SLR of 9 primary studies: four assessment dimensions (acceptance, ethics/security, skills, decision process) and two impacts (practices, roles). Calls for field, case and longitudinal studies."),
 ("A10", "Review of 91 articles: digital technologies transform information and structure, automate basic tasks and decisions, and reshape roles (job elimination, upskilling, deskilling, reskilling) and boundaries. Abstract only in the sheet."),
 ("A31", "Review: actual use of machine learning in management accounting is still very limited in the academic literature; competences in analytics, statistics and ML are needed."),
 ("A50", "Bibliometric (183 articles): the role moves from producing figures to interpreting data and supporting decisions; technology skills indispensable; COVID accelerated digitalisation."),
 ("A21", "Conceptual: AI shifts management control from deductive to inductive logics; reconfigures forms, practices and infrastructures; accounting as problematisation rather than 'answer machine'."),
 ("A44", "Methodological essay: digital trace data and non-linear strategy-structure-IS relations undermine contingency-style research designs."),
 ("A6",  "Historical review of AI in accounting; no findings in the sheet."),
]
sec_rows = [[a, ART[a]["short"], t] for a, t in SEC]

# ---------------------------------------------------------------- compose markdown
L = []
w = L.append
w("# AI / digitalisation and Management Control Systems: findings-first synthesis")
w("")
w(f"Source: `Lecture_Articles-2_copy.xlsx`, sheet *Fiches de Lecture* (50 rows, A1-A50). Everything below is taken from the reading sheet. Nothing was added from memory, and no article was re-read in full. Statistics are as reported in the sheet.")
w("")
w("## 0. Read this first: what the corpus can and cannot support")
w("")
w(f"- **Corpus composition.** {n_quant} quantitative or mixed studies, {n_qual_ok} qualitative studies with extractable findings, 2 qualitative studies with no extractable findings (A33, A42), 9 reviews/conceptual papers used only as corroboration, and 1 excluded vendor-video study (A36). That gives **{len(F)} coded empirical findings from {len({x['art'] for x in F})} articles.**")
w("- **The sheet's title says '38 articles' but holds 50 rows.** The notes on A42, A45, A47 and A50 refer to selection criteria (C4, C6) used to fill gaps in the first 38, so the corpus is a purposive selection tied to your framework, not a sample of the literature. Counts below describe this corpus.")
w("- **Row-level data problems that affect the analysis:**")
w("  - **A4**: the methodology, theory, variables and hypotheses cells are copied from A15 (an AI-transparency experiment). I used only A4's abstract and keywords (BI action research in an engineering group).")
w("  - **A33**: the 'conclusions' cell contains literature-review sections, not results. No finding could be extracted.")
w("  - **A42**: lists variables only ('lecture integrale requise'). No finding could be extracted.")
w("  - **A8, A17, A25, A35, A41, A29, A10, A6** have little more than an abstract. Findings from A8, A25, A35 and A41 are coded from abstracts and flagged. A17 and A41 give no effect sizes and are counted as 'unverified'.")
w("  - **A9** is labelled 'quantitative' in the sheet but is a 14-interview qualitative study. **A16** is labelled hypothetico-deductive but is interview-based. I coded the method from the description.")
w("  - **A5, A14, A23, A26** report significance or mediation but no coefficients in the sheet.")
w("- **Coding conventions.** 1 = the block's construct is itself in the finding (as variable, mechanism, mediator, moderator or outcome); 0.5 = indirectly present (proxy, or implied but not measured); 0 = absent. I did not code a block because the authors mention it in the literature review. Two judgement calls: (i) *information quality* is coded 0.5 on Sensemaking because the figure places it in that block, although it is a measure of informational input rather than interpretation; (ii) *competences* are coded under Conditions when a study uses them as an antecedent. The **Sensing** column is 1 in almost every finding because digitalisation is the corpus's selection criterion; it is context, not a discriminating variable.")
w("- **Relationship counts are computed, not estimated.** Each finding is tagged to the relationships it speaks to (stance: supports / indirect / null / mixed; mode: quantitative / descriptive / qualitative / abstract-only). The tables in section 3 are generated from those tags, so every cell traces to a finding ID. Where one article appears under both 'supporting' and 'null' it is because it contains both results (for example A2, A26, A47).")
w("- **Evidence strength.** Almost all quantitative studies are cross-sectional, self-reported and perceptual. The exceptions are the experiment A15 and the longitudinal job-advertisement study A47. Six of the 22 quantitative studies draw on German or DACH samples (A5, A7, A19, A26, A27, A46).")
w("")
w("## 1. Corpus triage")
w("")
w(table(["ID", "Short title", "Year", "Design", "Method", "Sample", "Context", "Use in this review", "Caveat"], tri_rows))
w("")
w("## 2. Empirical findings (one row per important result)")
w("")
w("Direction: + positive, - negative, 0 null, 'n/a' = no directional test. Significance is as reported; for qualitative studies it is not applicable.")
w("")
w(table(find_hdr, find_rows))
w("")
w("### 2b. Classification against the bibliometric structure")
w("")
w("1 = directly related; 0.5 = indirectly related; 0 = unrelated.")
w("")
w(table(code_hdr, code_rows))
w("")
w("### 2c. How much empirical weight does each block carry?")
w("")
w(table(bst_hdr, bst_rows))
w("")
w("Reading: Sensemaking has no quantitative article with a finding coded 1, Rewards has none at all, and Outcomes appear as a coded 1 only in quantitative studies. Planning has more empirical content here (4 quantitative, 1 qualitative articles with a finding coded 1) than the bibliometric n=1 suggests.")
w("")
w("## 3. Relationships that emerge empirically")
w("")
w("### 3a. Relationships you listed")
w("")
w(table(rel_hdr, req_rows))
w("")
w("### 3b. Relationships that the findings add (not in your list)")
w("")
w(table(rel_hdr, em_rows))
w("")
w("### 3c. Conditions")
w("")
w(table(rel_hdr, cond_rows))
w("")
w("## 4. Empirical puzzles")
w("")
w(table(["Empirical puzzle", "Supporting studies", "What the studies observe", "What remains unexplained"], [list(p) for p in N.PUZZLES]))
w("")
w("## 5. Missing mechanisms (derived from the puzzles only)")
w("")
w(table(["Proposed mechanism", "Empirical findings creating the need for it", "Direct evidence", "Indirect evidence", "Qualitative evidence", "Still untested?"], [list(m) for m in N.MECH]))
w("")
w("## 6. The proposed chain, link by link")
w("")
w("Digital capabilities -> informational capacity -> controller / business partner -> interaction -> interpretation / sensemaking -> managerial decision -> organisational action -> performance")
w("")
w("A = directly tested quantitatively; B = statistically associated, mechanism not tested; C = qualitatively documented; D = theoretically inferred; E = no evidence found; F = contradictory/null evidence exists. Several letters are given where a link has several kinds of evidence; the first is the strongest.")
w("")
w(table(["Link", "Evidence level", "Articles", "Main empirical finding", "Remaining gap"], [list(c) for c in N.CHAIN]))
w("")
w("### 6b. Is the literature fragmented?")
w(N.FRAGMENT_TEXT)
w("")
w("## 7. Synthesis by block of the conceptual structure")
w("")
for title, body in N.SYNTH:
    w(f"### {title}")
    w(body.strip())
    w("")
w("## 8. Answers to your ten questions")
w("")
for q, a in N.QA:
    w(f"**{q}**")
    w("")
    w(a)
    w("")
w("## Appendix A. Review-level corroboration (not counted as primary evidence)")
w("")
w(table(["ID", "Short title", "What the review says (from the sheet)"], sec_rows))
w("")
w("## Appendix B. A36 (excluded)")
w("")
w("A36 analyses three vendor promotional videos (ERP, RFID, big data). It reports that digitalisation 'strengthens efficiency and reduces uncertainty'. That is vendor discourse, not organisational evidence, so it is not counted.")
w("")
with open(os.path.join(OUT, "REPORT.md"), "w", encoding="utf-8") as fh:
    fh.write("\n".join(L))

# ---------------------------------------------------------------- xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()
def sheet(name, header, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.append(header)
    for r in rows:
        ws.append([str(c) if isinstance(c, (list, dict)) else c for c in r])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 18)
    ws.freeze_panes = "C2"
    return ws

sheet("Corpus", ["ID", "Short title", "Year", "Design", "Method", "Sample", "Context", "Use", "Caveat"], tri_rows,
      {"Short title": 40, "Method": 40, "Caveat": 50, "Context": 28, "Sample": 24, "Use": 24}, first=True)
sheet("Findings", find_hdr, find_rows, {"Empirical finding": 70, "Independent variable / antecedent": 30, "Mechanism": 28, "Outcome": 28,
      "Significant?": 24, "Method": 32, "Sample": 22, "Context": 26})
ws = sheet("Coding", code_hdr, code_rows, {"Empirical finding (short)": 60})
for row in ws.iter_rows(min_row=2, min_col=4):
    for c in row:
        c.alignment = Alignment(horizontal="center")
        if c.value == "1":
            c.fill = PatternFill("solid", fgColor="9BC2E6")
        elif c.value == "0.5":
            c.fill = PatternFill("solid", fgColor="DDEBF7")
sheet("BlockWeight", bst_hdr, bst_rows, {"Quantitative articles with a finding coded 1": 60, "Qualitative articles with a finding coded 1": 60})
wd = {"Relationship": 36, "Supporting empirical studies": 34, "Quantitative evidence": 60, "Qualitative evidence": 50, "Null / contradictory evidence": 50, "Reading": 70, "Number of studies": 26}
sheet("Relationships", rel_hdr, req_rows + em_rows, wd)
sheet("Conditions", rel_hdr, cond_rows, wd)
sheet("Puzzles", ["Empirical puzzle", "Supporting studies", "What the studies observe", "What remains unexplained"], [list(p) for p in N.PUZZLES],
      {"Empirical puzzle": 36, "Supporting studies": 34, "What the studies observe": 80, "What remains unexplained": 60})
sheet("Mechanisms", ["Proposed mechanism", "Findings creating the need", "Direct evidence", "Indirect evidence", "Qualitative evidence", "Still untested?"],
      [list(m) for m in N.MECH], {"Proposed mechanism": 34, "Findings creating the need": 50, "Direct evidence": 50, "Indirect evidence": 50, "Qualitative evidence": 40, "Still untested?": 40})
sheet("Chain", ["Link", "Evidence level", "Articles", "Main empirical finding", "Remaining gap"], [list(c) for c in N.CHAIN],
      {"Link": 34, "Evidence level": 18, "Articles": 40, "Main empirical finding": 80, "Remaining gap": 60})
sheet("Reviews", ["ID", "Short title", "What the review says"], sec_rows, {"Short title": 40, "What the review says": 120})
wb.save(os.path.join(OUT, "evidence_tables.xlsx"))
print("ok", len(L), "lines")
