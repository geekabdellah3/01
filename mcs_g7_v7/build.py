import re,subprocess,sys
import docx
D=docx.Document("/root/.claude/uploads/b8693ac0-eab9-5be8-ab8b-396dfef9021f/f943a14f-MCS_Practices_Transferability_WBES_G7_v6.docx")
def pipe(i):
    t=D.tables[i]; rows=[[c.text.replace("|","\\|").replace("\n"," ").strip() for c in r.cells] for r in t.rows]
    out=["| "+" | ".join(rows[0])+" |","|"+"|".join([":--"]+["--:"]*(len(rows[0])-1))+"|"]
    out+=["| "+" | ".join(r)+" |" for r in rows[1:]]
    return "\n".join(out)
M={(109,125):0,(131,143):1,(149,161):2,(167,187):3,(195,205):4,(209,225):5,(233,247):6,(253,269):7,(271,287):8,(295,317):9}
S="/tmp/claude-0/-home-user-01/b8693ac0-eab9-5be8-ab8b-396dfef9021f/scratchpad"
v=open(S+"/v6.md").read().split("\n")
b=open(sys.argv[1]).read()
def r(m):
    return pipe(M[(int(m[1]),int(m[2]))])
b=b.replace("x","x")
b=re.sub(r"@@T:(\d+)-(\d+)@@",r,b)
b=b.replace("Weighted WLS (wmedian)","Weighted WLS (sampling weights)").replace(">=10%","≥ 10%")
open("v7.md","w").write(b)
abs_=re.search(r"abstract: \|\n(.*?)\n---",b,re.S).group(1)
print("abstract words",len(abs_.split()))
subprocess.run(["pandoc","v7.md","-f","markdown","-o",sys.argv[2],"--resource-path=."],check=True)

import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
p=sys.argv[2]
d=docx.Document(p)
for st in d.styles:
    if st.type==1 and st.name in("Normal","Body Text","First Paragraph","Compact","Block Text"):
        st.font.name="Times New Roman"; st.font.size=Pt(11)
for t in d.tables:
    bd=OxmlElement("w:tblBorders")
    for e in("top","bottom","insideH"):
        x=OxmlElement("w:"+e); x.set(qn("w:val"),"single"); x.set(qn("w:sz"),"4"); x.set(qn("w:color"),"808080"); bd.append(x)
    t._tbl.tblPr.append(bd)
    for r in t.rows:
        for c in r.cells:
            for para in c.paragraphs:
                for run in para.runs: run.font.size=Pt(8.5)
d.save(p)
